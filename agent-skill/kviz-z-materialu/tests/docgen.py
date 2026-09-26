"""Programmatic generation of small test documents (DOCX, PPTX, PDF)."""
from __future__ import annotations

from pathlib import Path


def make_docx(path: Path) -> Path:
    from docx import Document

    doc = Document()
    doc.add_heading("Fotosyntéza", level=1)
    doc.add_paragraph("Fotosyntéza probíhá v chloroplastech zelených rostlin.")
    doc.add_paragraph("Rostlina přitom spotřebovává oxid uhličitý a vodu.")
    doc.add_heading("Dýchání", level=2)
    doc.add_paragraph("Při dýchání rostlina spotřebovává kyslík.")
    t = doc.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Proces", "Produkt"
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = "Fotosyntéza", "Glukóza"
    doc.save(str(path))
    return path


def make_pptx(path: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    layout = prs.slide_layouts[1]  # title + content
    slides = [
        ("Úvod", "Krátký úvod.", "Poznámka k úvodu."),
        ("Cíle", "Naučíme se základy.", ""),
        ("Buňka", "Buňka je základní stavební a funkční jednotka všech živých organismů. " * 4, "Řečník zdůrazní, že jádro obsahuje DNA."),
    ]
    for title, body, notes in slides:
        s = prs.slides.add_slide(layout)
        s.shapes.title.text = title
        s.placeholders[1].text = body
        if notes:
            s.notes_slide.notes_text_frame.text = notes
    s.shapes.add_textbox(Inches(1), Inches(6), Inches(3), Inches(1)).text_frame.text = "Doplňkový text."
    prs.save(str(path))
    return path


def _escape(s: str) -> str:
    return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(path: Path, pages: list[list[str]]) -> Path:
    """Minimal valid PDF with Helvetica text lines (ASCII only)."""
    objs: list[bytes] = []
    n_pages = len(pages)
    font_id = 3 + 2 * n_pages
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(n_pages))
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>".encode())
    for i, lines in enumerate(pages):
        content_id = 4 + 2 * i
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents {content_id} 0 R /Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode())
        ops = ["BT", "/F1 11 Tf", "14 TL", "50 800 Td"]
        for line in lines:
            ops.append(f"({_escape(line)}) Tj T*")
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1")
        objs.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))
    return path
