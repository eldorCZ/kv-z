#!/usr/bin/env python3
"""Generates fixtures/optika.docx – the sample teaching material used by the end-to-end test and quiz.json."""
from pathlib import Path

from docx import Document

SECTIONS = [
    (
        "Světlo a jeho šíření",
        [
            "Světlo je elektromagnetické vlnění, které vnímáme zrakem. Zdrojem světla může být Slunce, žárovka nebo plamen svíčky.",
            "Ve vakuu se světlo šíří rychlostí přibližně 300 000 km/s. V ostatních prostředích je jeho rychlost menší.",
            "Průhledná prostředí, například vzduch, voda nebo sklo, světlo propouštějí. Neprůhledná prostředí, jako dřevo nebo kov, světlo nepropouštějí.",
        ],
    ),
    (
        "Lom světla",
        [
            "Změnu směru šíření světla na rozhraní dvou prostředí nazýváme lom světla.",
            "Při přechodu do opticky hustšího prostředí se světlo láme ke kolmici. Při přechodu do opticky řidšího prostředí se světlo láme od kolmice.",
            "Index lomu vody je přibližně 1,33. Index lomu skla je asi 1,5 a diamantu 2,42.",
        ],
    ),
    (
        "Čočky",
        [
            "Spojka je čočka, která je uprostřed silnější než na okrajích. Rozptylka je naopak uprostřed tenčí.",
            "Rovnoběžné paprsky se po průchodu spojnou čočkou sbíhají v ohnisku. Vzdálenost ohniska od středu čočky se nazývá ohnisková vzdálenost.",
            "Spojky se používají v lupách, fotoaparátech a brýlích pro dalekozraké.",
        ],
    ),
]


def main() -> None:
    doc = Document()
    doc.add_heading("Optika pro 8. ročník", level=0)
    for title, paragraphs in SECTIONS:
        doc.add_heading(title, level=1)
        for p in paragraphs:
            doc.add_paragraph(p)
    doc.add_heading("Přehled indexů lomu", level=1)
    table = doc.add_table(rows=4, cols=2)
    for row, (a, b) in zip(table.rows, [("Látka", "Index lomu"), ("Vzduch", "1,00"), ("Voda", "1,33"), ("Diamant", "2,42")]):
        row.cells[0].text, row.cells[1].text = a, b
    out = Path(__file__).with_name("optika.docx")
    doc.save(str(out))
    print(out)


if __name__ == "__main__":
    main()
